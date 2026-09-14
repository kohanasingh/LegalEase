"""
python-docx-based text extraction for Word documents (.docx only — the
legacy binary .doc format needs external tools like LibreOffice/antiword
and isn't supported). Mirrors pdf_parser_tool.py's shape: a core
bytes-in-text-out function plus a CrewAI tool wrapper for agent use.
"""
import docx
from crewai.tools import tool


def extract_text_from_docx_bytes(data: bytes) -> str:
    import io

    document = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


@tool("DocxParserTool")
def docx_parser_tool(file_path: str) -> str:
    """Extracts plain text from a Word (.docx) file at the given path."""
    with open(file_path, "rb") as f:
        return extract_text_from_docx_bytes(f.read())
